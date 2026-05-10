from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


BASE_DIR = Path(__file__).resolve().parent.parent
MARKDOWN_FILE = BASE_DIR / "PAPER_DRAFT.md"
FIGURE_FILE = BASE_DIR / "figures" / "wifi7_vs_wifi6_summary.png"
OUTPUT_FILE = BASE_DIR / "WiFi7_vs_WiFi6_Final_Paper.docx"


def clean_inline_markdown(text: str) -> str:
    text = text.replace("`", "")
    text = text.replace("**", "")
    text = text.replace("\\[", "").replace("\\]", "")
    text = text.replace("\\times", "x")
    text = text.replace("\\frac{", "")
    text = text.replace("}{", " / ")
    text = text.replace("}", "")
    text = text.replace("\\approx", "approx.")
    return text.strip()


def add_title(document: Document, title: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(16)


def add_figure(document: Document) -> None:
    if not FIGURE_FILE.exists():
        return

    document.add_heading("Figure", level=1)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIGURE_FILE), width=Inches(6.3))

    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run("Figure 1. Summary comparison of Wi-Fi 7 and Wi-Fi 6 across the three experiments.").italic = True


def add_code_block(document: Document, code_lines: list[str]) -> None:
    for raw_line in code_lines:
        p = document.add_paragraph()
        run = p.add_run(raw_line.rstrip("\n"))
        run.font.name = "Courier New"
        run.font.size = Pt(9)


def main() -> None:
    if not MARKDOWN_FILE.exists():
        raise FileNotFoundError(f"Missing markdown draft: {MARKDOWN_FILE}")

    output_file = Path(sys.argv[1]) if len(sys.argv) > 1 else OUTPUT_FILE

    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)

    lines = MARKDOWN_FILE.read_text(encoding="utf-8").splitlines()
    inserted_figure = False

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "## 6. Limitations" and not inserted_figure:
            add_figure(document)
            inserted_figure = True

        if stripped.startswith("# "):
            add_title(document, clean_inline_markdown(stripped[2:]))
            i += 1
            continue

        if stripped.startswith("```"):
            i += 1
            code_lines: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code_block(document, code_lines)
            if i < len(lines) and lines[i].strip().startswith("```"):
                i += 1
            continue

        if stripped.startswith("## "):
            document.add_heading(clean_inline_markdown(stripped[3:]), level=1)
            i += 1
            continue

        if stripped.startswith("### "):
            document.add_heading(clean_inline_markdown(stripped[4:]), level=2)
            i += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            document.add_paragraph(clean_inline_markdown(re.sub(r"^\d+\.\s*", "", stripped)), style="List Number")
            i += 1
            continue

        if stripped.startswith("- "):
            document.add_paragraph(clean_inline_markdown(stripped[2:]), style="List Bullet")
            i += 1
            continue

        paragraph_lines = [clean_inline_markdown(stripped)]
        i += 1
        while i < len(lines):
            next_line = lines[i].strip()
            if not next_line:
                break
            if (
                next_line.startswith("#")
                or next_line.startswith("- ")
                or next_line.startswith("```")
                or re.match(r"^\d+\.\s", next_line)
            ):
                break
            paragraph_lines.append(clean_inline_markdown(next_line))
            i += 1

        paragraph_text = " ".join(part for part in paragraph_lines if part)
        if paragraph_text:
            document.add_paragraph(paragraph_text)

    if not inserted_figure:
        add_figure(document)

    document.save(str(output_file))
    print(f"Saved DOCX: {output_file}")


if __name__ == "__main__":
    main()
